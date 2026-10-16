"""Planilhas: exportação de pedidos e catálogo (CSV/XLSX) e leitura do catálogo para importação.

CSV sai no padrão brasileiro (";" como separador, vírgula decimal, UTF-8 com BOM) para abrir
direto no LibreOffice ou no Excel.
"""

import csv
import io
from datetime import datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from app.config import Settings
from app.models import Order, Product
from app.services.text import normalize

ORDER_HEADER = [
    "data",
    "pedido",
    "status",
    "cliente",
    "telefone",
    "sku",
    "item",
    "quantidade",
    "unidade",
    "preco_unitario",
    "subtotal",
    "total_itens",
    "frete",
    "total_pedido",
    "modalidade",
    "endereco_entrega",
    "pagamento",
    "troco_para",
    "pago_em",
]

CATALOG_HEADER = [
    "sku",
    "nome_tecnico",
    "categoria",
    "unidade_venda",
    "preco",
    "atributos",
    "sinonimos",
    "ativo",
]

MAX_IMPORT_BYTES = 5 * 1024 * 1024


class ImportFileError(Exception):
    pass


def _local(dt: datetime | None, settings: Settings) -> datetime | None:
    return dt.astimezone(settings.tz).replace(tzinfo=None) if dt else None


def order_rows(orders: list[Order], settings: Settings) -> list[list[Any]]:
    """Uma linha por item, com os dados do pedido repetidos (fácil de filtrar e somar)."""
    rows: list[list[Any]] = []
    for order in orders:
        customer = order.customer
        for item in order.items:
            preco = item.preco_unitario or Decimal(0)
            rows.append(
                [
                    _local(order.confirmado_em or order.created_at, settings),
                    order.numero,
                    order.status.value,
                    (customer.nome or customer.nome_whatsapp or "") if customer else "",
                    (customer.telefone_e164 or "") if customer else "",
                    item.product.sku if item.product else "",
                    item.item_normalizado,
                    item.quantidade,
                    item.unidade,
                    preco,
                    preco * item.quantidade,
                    order.total,
                    order.frete,
                    order.total_final,
                    order.modalidade.value if order.modalidade else "",
                    order.endereco_entrega or "",
                    order.forma_pagamento.value if order.forma_pagamento else "",
                    order.troco_para or "",
                    _local(order.pago_em, settings),
                ]
            )
    return rows


def product_rows(products: list[Product]) -> list[list[Any]]:
    return [
        [
            p.sku,
            p.nome_tecnico,
            p.categoria or "",
            p.unidade_venda,
            p.preco,
            "; ".join(f"{k}={v}" for k, v in (p.atributos or {}).items()),
            "; ".join(s.termo for s in p.synonyms),
            "sim" if p.ativo else "não",
        ]
        for p in products
    ]


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y %H:%M")
    if isinstance(value, Decimal):
        return f"{value.normalize():f}".replace(".", ",")
    return str(value)


def to_csv(header: list[str], rows: list[list[Any]]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([_csv_value(v) for v in row])
    return buffer.getvalue().encode("utf-8-sig")


def to_xlsx(header: list[str], rows: list[list[Any]], title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for col, name in enumerate(header, 1):
        letter = ws.cell(row=1, column=col).column_letter
        values = [len(name)] + [len(_csv_value(r[col - 1])) for r in rows[:200]]
        ws.column_dimensions[letter].width = min(max(values) + 2, 60)
        if name in ("data", "pago_em"):
            for cell in ws[letter][1:]:
                cell.number_format = "dd/mm/yyyy hh:mm"
        elif name.startswith(("preco", "subtotal", "total", "frete")):
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _key(name: Any) -> str:
    return normalize(str(name or "")).replace(" ", "_")


def read_table(filename: str, data: bytes) -> list[dict[str, Any]]:
    """Lê um CSV (; ou ,) ou XLSX com cabeçalho na primeira linha."""
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportFileError("Arquivo maior que 5 MB")
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise ImportFileError("Não foi possível ler o arquivo XLSX") from exc
        rows = list(wb.active.iter_rows(values_only=True))
        if not rows:
            return []
        header = [_key(h) for h in rows[0]]
        return [
            {h: ("" if v is None else v) for h, v in zip(header, row, strict=False) if h}
            for row in rows[1:]
            if any(v not in (None, "") for v in row)
        ]
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        first_line = text.splitlines()[0] if text else ""
        delimiter = ";" if first_line.count(";") > first_line.count(",") else ","
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        return [{_key(k): v for k, v in row.items() if k} for row in reader]
    raise ImportFileError("Envie um arquivo .csv ou .xlsx")
