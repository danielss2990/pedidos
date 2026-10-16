import hashlib
import hmac
import io
import json
from datetime import UTC, datetime, timedelta
from decimal import Decimal
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.main import create_app
from app.models import (
    Conversation,
    ConversationState,
    Customer,
    FormaPagamento,
    Message,
    MessageDirection,
    MessageType,
    Modalidade,
    Order,
    OrderItem,
    OrderStatus,
    Role,
    SentBy,
    User,
)
from app.security import hash_password
from tests.conftest import FakeMessenger


class FakeQueue:
    def __init__(self) -> None:
        self.jobs: list[tuple[str, tuple[Any, ...], dict[str, Any]]] = []

    async def enqueue_job(self, name: str, *args: Any, **kwargs: Any) -> None:
        self.jobs.append((name, args, kwargs))


@pytest.fixture
async def api():
    app = create_app(use_lifespan=False)
    app.state.messenger = FakeMessenger()
    app.state.queue = FakeQueue()
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        client.app = app
        yield client


async def _user(session, perfil: Role, email: str) -> User:
    user = User(
        nome=email,
        email=email,
        senha_hash=hash_password("senha-forte-1"),
        perfil=perfil,
        ativo=True,
    )
    session.add(user)
    await session.commit()
    return user


async def _login(api, email: str) -> dict[str, str]:
    r = await api.post("/api/v1/auth/login", json={"email": email, "senha": "senha-forte-1"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _sign(body: bytes) -> str:
    return "sha256=" + hmac.new(b"segredo-meta", body, hashlib.sha256).hexdigest()


async def test_verificacao_do_webhook(api):
    params = {
        "hub.mode": "subscribe",
        "hub.verify_token": "token-verificacao",
        "hub.challenge": "42",
    }
    r = await api.get("/webhooks/whatsapp", params=params)
    assert r.status_code == 200 and r.text == "42"
    params["hub.verify_token"] = "errado"
    assert (await api.get("/webhooks/whatsapp", params=params)).status_code == 403


async def test_webhook_enfileira_mensagens(api):
    payload = {
        "entry": [
            {
                "changes": [
                    {
                        "field": "messages",
                        "value": {
                            "messages": [
                                {
                                    "from": "5511",
                                    "id": "wamid.X",
                                    "timestamp": "1",
                                    "type": "text",
                                    "text": {"body": "oi"},
                                }
                            ]
                        },
                    }
                ]
            }
        ]
    }
    body = json.dumps(payload).encode()
    r = await api.post(
        "/webhooks/whatsapp", content=body, headers={"X-Hub-Signature-256": _sign(body)}
    )
    assert r.status_code == 200 and r.json() == {"recebidas": 1}
    name, args, kwargs = api.app.state.queue.jobs[0]
    assert name == "process_incoming"
    assert args[0]["text"] == "oi"
    assert kwargs["_job_id"] == "wa:wamid.X"


async def test_webhook_rejeita_assinatura_invalida(api):
    r = await api.post(
        "/webhooks/whatsapp", content=b"{}", headers={"X-Hub-Signature-256": "sha256=00"}
    )
    assert r.status_code == 401
    assert api.app.state.queue.jobs == []


async def test_login_e_perfis(api, session):
    await _user(session, Role.ADMIN, "admin@loja.com")
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    admin = await _login(api, "admin@loja.com")
    atendente = await _login(api, "ana@loja.com")

    assert (await api.get("/api/v1/auth/me", headers=atendente)).json()["perfil"] == "atendente"
    assert (await api.get("/api/v1/users", headers=atendente)).status_code == 403
    assert (await api.get("/api/v1/users", headers=admin)).status_code == 200
    assert (await api.get("/api/v1/users")).status_code == 401

    bad = await api.post("/api/v1/auth/login", json={"email": "ana@loja.com", "senha": "errada!!"})
    assert bad.status_code == 401

    r = await api.post(
        "/api/v1/users",
        headers=admin,
        json={
            "nome": "Sueli",
            "email": "sueli@loja.com",
            "senha": "outra-senha-1",
            "perfil": "supervisor",
        },
    )
    assert r.status_code == 201


async def test_refresh_por_cookie(api, session):
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    await _login(api, "ana@loja.com")
    r = await api.post("/api/v1/auth/refresh")
    assert r.status_code == 200 and r.json()["access_token"]


async def test_atendente_assume_responde_e_devolve(api, session):
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    headers = await _login(api, "ana@loja.com")
    customer = Customer(telefone_e164="5511999998888")
    session.add(customer)
    await session.flush()
    conversation = Conversation(
        customer_id=customer.id,
        estado=ConversationState.ATENDIMENTO_HUMANO,
        duvidas_sem_solucao=0,
        ultimo_contato_cliente_em=datetime.now(UTC),
    )
    session.add(conversation)
    await session.commit()
    base = f"/api/v1/conversations/{conversation.id}"

    session.add(
        Message(
            conversation_id=conversation.id,
            direcao=MessageDirection.ENTRADA,
            tipo=MessageType.AUDIO,
            conteudo=None,
            transcricao="quero falar com alguém",
            enviado_por=SentBy.CLIENTE,
        )
    )
    await session.commit()
    lista = await api.get("/api/v1/conversations?estado=atendimento_humano", headers=headers)
    assert lista.json()["total"] == 1
    assert lista.json()["items"][0]["ultima_mensagem"]["conteudo"] == "quero falar com alguém"

    assert (
        await api.post(f"{base}/messages", headers=headers, json={"texto": "Oi!"})
    ).status_code == 409
    detalhe = await api.get(base, headers=headers)
    assert detalhe.json()["customer"]["telefone_e164"] == "5511999998888"
    assumida = await api.post(f"{base}/assume", headers=headers)
    assert assumida.json()["atendente"]["nome"] == "ana@loja.com"
    r = await api.post(f"{base}/messages", headers=headers, json={"texto": "Oi, sou a Ana."})
    assert r.status_code == 201 and r.json()["enviado_por"] == "atendente"
    assert api.app.state.messenger.last == "Oi, sou a Ana."

    r = await api.post(f"{base}/release", headers=headers)
    assert r.json()["estado"] == "coletando"
    msgs = await api.get(f"{base}/messages", headers=headers)
    assert [m["conteudo"] for m in msgs.json()] == [None, "Oi, sou a Ana."]


async def test_catalogo_crud_e_busca(api, session):
    await _user(session, Role.ADMIN, "admin@loja.com")
    admin = await _login(api, "admin@loja.com")
    r = await api.post(
        "/api/v1/products",
        headers=admin,
        json={
            "sku": "ABR-NY-10",
            "nome_tecnico": "Abraçadeira de Nylon 10cm",
            "preco": "12.50",
            "sinonimos": ["enforca gato", "bucha"],
        },
    )
    assert r.status_code == 201
    busca = await api.get("/api/v1/products?q=enforca", headers=admin)
    assert busca.json()["items"][0]["sinonimos"] == ["enforca gato", "bucha"]

    pid = r.json()["id"]
    r = await api.patch(f"/api/v1/products/{pid}", headers=admin, json={"preco": "13.00"})
    assert r.json()["preco"] == "13.00"


async def test_catalogo_importa_csv_e_xlsx_e_exporta(api, session):
    await _user(session, Role.ADMIN, "admin@loja.com")
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    admin = await _login(api, "admin@loja.com")
    atendente = await _login(api, "ana@loja.com")

    # CSV no padrão do Excel/LibreOffice em português: ";" e vírgula decimal.
    csv_body = (
        "SKU;Nome Técnico;Categoria;Unidade Venda;Preço;Atributos;Sinônimos;Ativo\r\n"
        "ABR-10;Abraçadeira de Nylon 10cm;acessorios;pacotes;12,50;;enforca gato;sim\r\n"
        "SEM-PRECO;Item sem preço;;;;;;sim\r\n"
    ).encode("utf-8-sig")
    files = {"arquivo": ("catalogo.csv", csv_body, "text/csv")}
    r = await api.post("/api/v1/products/import", headers=admin, files=files)
    assert r.json() == {"criados": 1, "atualizados": 0, "ignorados": 1}
    negado = await api.post("/api/v1/products/import", headers=atendente, files=files)
    assert negado.status_code == 403

    exportado = await api.get("/api/v1/products/export?formato=xlsx", headers=admin)
    assert exportado.headers["content-disposition"].endswith('catalogo.xlsx"')
    # O XLSX exportado volta pela importação (o cliente edita e reenvia).
    files = {"arquivo": ("catalogo.xlsx", exportado.content, "application/octet-stream")}
    r = await api.post("/api/v1/products/import", headers=admin, files=files)
    assert r.json() == {"criados": 0, "atualizados": 1, "ignorados": 0}

    ruim = {"arquivo": ("lista.txt", b"oi", "text/plain")}
    r = await api.post("/api/v1/products/import", headers=admin, files=ruim)
    assert r.status_code == 422
    sem_colunas = {"arquivo": ("x.csv", b"codigo;descricao\n1;a\n", "text/csv")}
    r = await api.post("/api/v1/products/import", headers=admin, files=sem_colunas)
    assert r.status_code == 422 and "nome_tecnico" in r.json()["detail"]


async def test_horario_comercial_api(api, session):
    await _user(session, Role.ADMIN, "admin@loja.com")
    admin = await _login(api, "admin@loja.com")
    padrao = await api.get("/api/v1/settings/business-hours", headers=admin)
    assert len(padrao.json()) == 6
    novo = [{"dia_semana": 0, "abre": "07:30", "fecha": "17:00"}]
    r = await api.put("/api/v1/settings/business-hours", headers=admin, json=novo)
    assert r.status_code == 200
    assert (await api.get("/api/v1/settings/business-hours", headers=admin)).json()[0][
        "abre"
    ] == "07:30:00"


async def _pedido_fechado(session, ultimo_contato: datetime) -> tuple[Conversation, Order]:
    customer = Customer(telefone_e164="5511999998888")
    session.add(customer)
    await session.flush()
    conversation = Conversation(
        customer_id=customer.id,
        estado=ConversationState.PEDIDO_EM_ANDAMENTO,
        duvidas_sem_solucao=0,
        ultimo_contato_cliente_em=ultimo_contato,
    )
    session.add(conversation)
    await session.flush()
    order = Order(
        conversation_id=conversation.id,
        customer_id=customer.id,
        status=OrderStatus.CONFIRMADO,
        total=Decimal("100.00"),
        modalidade=Modalidade.ENTREGA,
        endereco_entrega="Rua A, 1",
        forma_pagamento=FormaPagamento.PIX,
        confirmado_em=datetime.now(UTC),
        finalizar_pendente=False,
    )
    session.add(order)
    await session.commit()
    return conversation, order


async def test_equipe_acompanha_pedido_ate_concluir(api, session):
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    headers = await _login(api, "ana@loja.com")
    conversation, order = await _pedido_fechado(session, datetime.now(UTC))
    base = f"/api/v1/orders/{order.id}"
    wa = api.app.state.messenger

    andamento = await api.get("/api/v1/orders?em_andamento=true", headers=headers)
    assert andamento.json()["total"] == 1

    r = await api.post(f"{base}/status", headers=headers, json={"status": "pronto"})
    assert r.status_code == 409 and "frete" in r.json()["detail"]

    r = await api.post(f"{base}/frete", headers=headers, json={"valor": "15.00"})
    body = r.json()
    assert body["mensagem_enviada"] is True
    assert body["order"]["total_final"] == "115.00"
    assert "Total a pagar: R$ 115,00" in wa.last and "pix@loja.com" in wa.last

    r = await api.post(f"{base}/pagamento", headers=headers, json={"pago": True})
    assert r.json()["order"]["pago_em"] is not None and r.json()["mensagem_enviada"] is False

    r = await api.post(f"{base}/status", headers=headers, json={"status": "em_separacao"})
    assert "está sendo separado" in wa.last
    r = await api.post(f"{base}/status", headers=headers, json={"status": "pronto"})
    assert "saiu para entrega" in wa.last
    r = await api.post(f"{base}/status", headers=headers, json={"status": "concluido"})
    assert r.json()["order"]["status"] == "concluido"
    assert "Obrigado pela preferência" in wa.last

    await session.refresh(conversation)
    assert conversation.estado == ConversationState.COLETANDO

    r = await api.post(f"{base}/status", headers=headers, json={"status": "pronto"})
    assert r.status_code == 409

    msgs = await api.get(f"/api/v1/conversations/{conversation.id}/messages", headers=headers)
    assert all(m["enviado_por"] == "atendente" for m in msgs.json())


async def test_aviso_fora_da_janela_de_24h(api, session):
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    headers = await _login(api, "ana@loja.com")
    _, order = await _pedido_fechado(session, datetime.now(UTC) - timedelta(days=2))
    r = await api.post(
        f"/api/v1/orders/{order.id}/status", headers=headers, json={"status": "cancelado"}
    )
    body = r.json()
    assert body["order"]["status"] == "cancelado"
    assert body["mensagem_enviada"] is False and "24 h" in body["aviso"]
    assert api.app.state.messenger.sent == []


async def test_exporta_pedidos_em_csv_e_xlsx(api, session):
    await _user(session, Role.ATENDENTE, "ana@loja.com")
    headers = await _login(api, "ana@loja.com")
    _, order = await _pedido_fechado(session, datetime.now(UTC))
    session.add(
        OrderItem(
            order_id=order.id,
            posicao=1,
            descricao_original="2 tomadas",
            item_normalizado="Tomada 10A",
            quantidade=Decimal("2"),
            unidade="unidades",
            preco_unitario=Decimal("8.90"),
            opcoes=[],
        )
    )
    order.frete = Decimal("15.00")
    await session.commit()

    r = await api.get("/api/v1/orders/export?formato=csv", headers=headers)
    text = r.content.decode("utf-8-sig")
    header, linha = text.strip().split("\r\n")
    assert header.startswith("data;pedido;status;cliente")
    assert ";Tomada 10A;2;unidades;8,9;17,8;100;15;115;entrega;Rua A, 1;pix;" in linha

    r = await api.get("/api/v1/orders/export?formato=xlsx", headers=headers)
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws["A1"].value == "data" and ws.max_row == 2
    assert ws["N2"].value == 115  # total_pedido numérico, não texto

    hoje = datetime.now(UTC).date()
    vazio = await api.get(
        f"/api/v1/orders/export?formato=csv&de={hoje + timedelta(days=2)}", headers=headers
    )
    assert vazio.content.decode("utf-8-sig").strip().count("\r\n") == 0
